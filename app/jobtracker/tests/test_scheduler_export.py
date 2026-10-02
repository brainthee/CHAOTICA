"""Tests for the global scheduler XLSX export (toolbar Export button)."""
import io
import json
from datetime import timedelta

from django.test import RequestFactory
from guardian.shortcuts import assign_perm
from openpyxl import load_workbook

from chaotica_utils.models import User
from jobtracker.enums import DefaultTimeSlotTypes, TimeSlotDeliveryRole
from jobtracker.models import (
    Client,
    Job,
    OrganisationalUnit,
    OrganisationalUnitMember,
    Phase,
    Project,
    TimeSlot,
    TimeSlotType,
)
from jobtracker.schedule_export import SCHEDULER_EXPORT_MAX_DAYS
from jobtracker.views import scheduler as sched_views
from .test_schedule_history import ScheduleHistoryBase


class SchedulerScopeBase(ScheduleHistoryBase):
    """actor + other share a unit the actor may view; hidden sits in another."""

    def setUp(self):
        super().setUp()
        self.rf = RequestFactory()
        # User.save() promotes the only user to superuser + Global Admin; on a
        # freshly flushed DB (just guardian's AnonymousUser) that's the actor.
        # Scope tests need an ordinary user, so demote explicitly.
        self.actor.is_superuser = self.actor.is_staff = False
        self.actor.first_name, self.actor.last_name = "Ada", "Actor"
        self.actor.save()
        self.actor.groups.clear()
        self.other.first_name, self.other.last_name = "Olly", "Other"
        self.other.save()
        OrganisationalUnitMember.objects.create(unit=self.unit, member=self.actor)
        OrganisationalUnitMember.objects.create(unit=self.unit, member=self.other)
        # Someone in a unit the actor can't see the schedule of.
        self.hidden = User.objects.create_user(
            email="hidden@test.com", password="pw12345", first_name="Hal", last_name="Hidden"
        )
        self.hidden_unit = OrganisationalUnit.objects.create(name="Hidden Unit")
        OrganisationalUnitMember.objects.create(unit=self.hidden_unit, member=self.hidden)
        assign_perm("jobtracker.view_users_schedule", self.actor, self.unit)
        self._reload_actor()

    def _reload_actor(self):
        self.actor = User.objects.get(pk=self.actor.pk)  # reset perm cache

    def _export(self, params=None, user=None, start=None, end=None):
        query = {
            "start": (start or self.start - timedelta(days=7)).isoformat(),
            "end": (end or self.end + timedelta(days=7)).isoformat(),
        }
        query.update(params or {})
        req = self.rf.get("/jobtracker/scheduler/export", query)
        req.user = user or self.actor
        return sched_views.export_scheduler_view(req)

    def _workbook(self, resp):
        self.assertEqual(resp.status_code, 200, resp.content[:200])
        self.assertIn("attachment;", resp["Content-Disposition"])
        return load_workbook(io.BytesIO(resp.content))

    def _resources(self, wb):
        ws = wb["Schedule"]
        return [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]


class SchedulerExportTests(SchedulerScopeBase):
    def test_export_has_all_sheets_and_visible_members_only(self):
        wb = self._workbook(self._export())
        self.assertEqual(wb.sheetnames, ["Overview", "Schedule", "Bookings", "Resources"])
        names = self._resources(wb)
        self.assertIn(self.actor.get_full_name(), names)
        self.assertIn(self.other.get_full_name(), names)  # nothing booked, still listed
        self.assertNotIn(self.hidden.get_full_name(), names)

    def test_grid_and_bookings_show_slot(self):
        slot = self._delivery_slot()
        wb = self._workbook(self._export())
        ws = wb["Schedule"]
        row = self._resources(wb).index(self.actor.get_full_name()) + 2
        cells = [c.value for c in ws[row] if c.value]
        self.assertTrue(any(slot.get_schedule_title() in str(v) for v in cells))
        bookings = wb["Bookings"]
        self.assertEqual(bookings.max_row, 2)
        self.assertEqual(bookings.cell(row=2, column=1).value, self.actor.get_full_name())
        self.assertEqual(bookings.cell(row=2, column=5).value, "{}: {}".format(self.job.id, self.job.title))

    def test_grid_spans_requested_range(self):
        start = self.start - timedelta(days=7)
        end = self.end + timedelta(days=7)
        ws = self._workbook(self._export(start=start, end=end))["Schedule"]
        meta_cols = 5
        days = (end.date() - start.date()).days + 1
        self.assertEqual(ws.max_column, meta_cols + days)

    def test_filter_is_applied(self):
        wb = self._workbook(self._export({"users": [str(self.other.pk)]}))
        self.assertEqual(self._resources(wb), [self.other.get_full_name()])
        overview = [c.value for row in wb["Overview"].iter_rows() for c in row if c.value]
        self.assertTrue(any(self.other.get_full_name() in str(v) for v in overview))

    def test_out_of_window_slot_excluded(self):
        self._delivery_slot()
        far = self.start + timedelta(days=60)
        wb = self._workbook(self._export(start=far, end=far + timedelta(days=7)))
        self.assertEqual(wb["Bookings"].max_row, 1)

    def test_missing_or_bad_range_is_400(self):
        req = self.rf.get("/jobtracker/scheduler/export", {"start": "nope"})
        req.user = self.actor
        self.assertEqual(sched_views.export_scheduler_view(req).status_code, 400)
        resp = self._export(start=self.end, end=self.start - timedelta(days=2))
        self.assertEqual(resp.status_code, 400)

    def test_range_cap(self):
        resp = self._export(end=self.start + timedelta(days=SCHEDULER_EXPORT_MAX_DAYS + 1))
        self.assertEqual(resp.status_code, 400)

    def test_members_feed_still_json(self):
        # Refactor guard: get_scheduler_members now wraps scheduler_member_rows.
        req = self.rf.get(
            "/jobtracker/scheduler/members",
            {"start": self.start.isoformat(), "end": self.end.isoformat()},
        )
        req.user = self.actor
        data = json.loads(sched_views.view_scheduler_members(req).content)
        self.assertEqual({r["id"] for r in data}, {self.actor.pk, self.other.pk})
        self.assertTrue(all(r["html_view"] for r in data))


class HiddenWorkBase(SchedulerScopeBase):
    """Adds a job, phase and project in the hidden unit, all booked for hidden."""

    def setUp(self):
        super().setUp()
        self.hidden_job = Job.objects.create(
            unit=self.hidden_unit, client=Client.objects.create(name="Hidden Client"),
            title="Hidden Job", created_by=self.hidden, account_manager=self.hidden,
        )
        self.hidden_phase = Phase.objects.create(job=self.hidden_job, title="Hidden Phase")
        TimeSlot.objects.create(
            user=self.hidden, slot_type=self.delivery_type, phase=self.hidden_phase,
            deliveryRole=TimeSlotDeliveryRole.DELIVERY, start=self.start, end=self.end,
        )
        self.hidden_project = Project.objects.create(
            title="Hidden Project", unit=self.hidden_unit, created_by=self.hidden
        )
        TimeSlot.objects.create(
            user=self.hidden, project=self.hidden_project, start=self.start, end=self.end,
            slot_type=TimeSlotType.get_builtin_object(DefaultTimeSlotTypes.INTERNAL_PROJECT),
        )

    def _member_ids(self, params, user=None):
        query = {"start": self.start.isoformat(), "end": self.end.isoformat()}
        query.update(params)
        req = self.rf.get("/jobtracker/scheduler/members", query)
        req.user = user or self.actor
        return {r["id"] for r in json.loads(sched_views.view_scheduler_members(req).content)}


class SchedulerFilterScopeTests(HiddenWorkBase):
    """Query-string filters must never widen the global scheduler beyond the
    viewer's view_users_schedule scope (shared by the feeds and the export)."""

    def test_include_user_cannot_add_out_of_scope_user(self):
        ids = self._member_ids({"include_user": [str(self.hidden.pk)]})
        self.assertNotIn(self.hidden.pk, ids)
        wb = self._workbook(self._export({"include_user": [str(self.hidden.pk)]}))
        self.assertNotIn(self.hidden.get_full_name(), self._resources(wb))

    def test_include_user_still_adds_in_scope_user(self):
        ids = self._member_ids({
            "users": [str(self.actor.pk)], "include_user": [str(self.other.pk)],
        })
        self.assertEqual(ids, {self.actor.pk, self.other.pk})

    def test_job_filter_on_unviewable_job_hides_team(self):
        ids = self._member_ids({"jobs": [str(self.hidden_job.pk)]})
        self.assertNotIn(self.hidden.pk, ids)

    def test_job_filter_on_viewable_job_shows_whole_team(self):
        assign_perm("jobtracker.view_job_schedule", self.actor, self.hidden_unit)
        self._reload_actor()
        ids = self._member_ids({"jobs": [str(self.hidden_job.pk)]})
        self.assertIn(self.hidden.pk, ids)

    def test_phase_filter_on_unviewable_phase_hides_team(self):
        ids = self._member_ids({"phases": [str(self.hidden_phase.pk)]})
        self.assertNotIn(self.hidden.pk, ids)

    def test_project_filter_respects_view_project(self):
        params = {"projects": [str(self.hidden_project.pk)]}
        self.assertNotIn(self.hidden.pk, self._member_ids(params))
        assign_perm("jobtracker.view_project", self.actor)
        self._reload_actor()
        self.assertIn(self.hidden.pk, self._member_ids(params))


class SchedulerClientFilterTests(HiddenWorkBase):
    """Clients filter: people booked on the client's jobs, other work faded."""

    def setUp(self):
        super().setUp()
        self._delivery_slot()  # actor on self.job (Test Client); other has nothing

    def _slots(self, params):
        query = {
            "start": (self.start - timedelta(days=1)).isoformat(),
            "end": (self.end + timedelta(days=1)).isoformat(),
        }
        query.update(params)
        req = self.rf.get("/jobtracker/scheduler/timeslots", query)
        req.user = self.actor
        return json.loads(sched_views.view_scheduler_slots(req).content)

    def _delivery_rows(self, params):
        return [r for r in self._slots(params) if r.get("phaseId") == self.phase.pk]

    def test_client_filter_shows_only_people_booked_for_client(self):
        ids = self._member_ids({"clients": [str(self.client_obj.pk)]})
        self.assertEqual(ids, {self.actor.pk})

    def test_client_filter_does_not_fade_its_own_work(self):
        rows = self._delivery_rows({"clients": [str(self.client_obj.pk)]})
        self.assertEqual(len(rows), 1)
        self.assertNotEqual(rows[0].get("display"), "background")

    def test_other_clients_work_is_faded(self):
        other_client = Client.objects.create(name="Elsewhere Ltd")
        rows = self._delivery_rows({
            "clients": [str(other_client.pk)], "include_user": [str(self.actor.pk)],
        })
        self.assertEqual(rows[0].get("display"), "background")

    def test_client_filter_respects_job_schedule_permission(self):
        params = {"clients": [str(self.hidden_job.client.pk)]}
        self.assertNotIn(self.hidden.pk, self._member_ids(params))
        assign_perm("jobtracker.view_job_schedule", self.actor, self.hidden_unit)
        self._reload_actor()
        self.assertIn(self.hidden.pk, self._member_ids(params))

    def test_deleted_jobs_are_ignored(self):
        from jobtracker.enums import JobStatuses

        Job.objects.filter(pk=self.job.pk).update(status=JobStatuses.DELETED)
        ids = self._member_ids({"clients": [str(self.client_obj.pk)]})
        self.assertNotIn(self.actor.pk, ids)

    def test_export_mutes_slots_outside_client(self):
        other_client = Client.objects.create(name="Elsewhere Ltd")
        wb = self._workbook(self._export({
            "clients": [str(other_client.pk)], "include_user": [str(self.actor.pk)],
        }))
        overview = " ".join(
            str(c.value) for row in wb["Overview"].iter_rows() for c in row if c.value
        )
        self.assertIn("Elsewhere Ltd", overview)
        ws = wb["Schedule"]
        row = self._resources(wb).index(self.actor.get_full_name()) + 2
        booked = [c for c in ws[row] if c.value and "Phase 1" in str(c.value)]
        self.assertTrue(booked)
        self.assertEqual(booked[0].fill.fgColor.rgb[-6:].lower(), "eff2f6")  # muted

    def test_form_offers_only_viewable_clients(self):
        from jobtracker.forms import SchedulerFilter

        form = SchedulerFilter({}, user=self.actor)
        self.assertNotIn(self.client_obj, form.fields["clients"].queryset)
        assign_perm("jobtracker.view_client", self.actor)
        self._reload_actor()
        form = SchedulerFilter({}, user=self.actor)
        self.assertIn(self.client_obj, form.fields["clients"].queryset)
